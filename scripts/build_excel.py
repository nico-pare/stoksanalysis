#!/usr/bin/env python3
"""Build a complete multi-sheet stock-analysis Excel workbook from the raw
JSONL data fetched by the worker agents (one record per symbol in data/out/).

Data source: Financial Modeling Prep (via MCP). Values are as-returned by the API.
"""
import glob
import json
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "data", "out")
SP500_FILE = os.path.join(ROOT, "data", "sp500_set.txt")
XLSX = os.path.join(ROOT, "Stock_Analysis_1000.xlsx")

SP500 = set(open(SP500_FILE).read().split())


def num(v):
    """Coerce to float when possible, else return the value unchanged/None."""
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return v


def pct(v):
    n = num(v)
    return n * 100 if isinstance(n, float) else None


def load_records():
    recs = {}
    for path in sorted(glob.glob(os.path.join(OUT_DIR, "*.jsonl"))):
        with open(path) as f:
            for line in f:
                line = line.strip().rstrip(",")
                if not line:
                    continue
                try:
                    obj = json.loads(line)
                except json.JSONDecodeError:
                    continue
                sym = obj.get("symbol")
                if sym:
                    recs[sym] = obj  # later chunks overwrite dupes
    return recs


def build_row(sym, obj):
    p = obj.get("profile") or {}
    r = obj.get("ratios") or {}
    t = obj.get("target") or {}
    g = obj.get("grades") or {}

    price = num(p.get("price"))
    tcons = num(t.get("targetConsensus"))
    upside = ((tcons / price - 1) * 100) if (price and tcons) else None

    sb = num(g.get("strongBuy")) or 0
    b = num(g.get("buy")) or 0
    h = num(g.get("hold")) or 0
    s = num(g.get("sell")) or 0
    ss = num(g.get("strongSell")) or 0
    total = sb + b + h + s + ss
    pct_buy = ((sb + b) / total * 100) if total else None
    pct_sbuy = (sb / total * 100) if total else None

    return {
        "Ticker": sym,
        "Company": p.get("companyName"),
        "In S&P 500": "Yes" if sym in SP500 else "No",
        "Sector": p.get("sector"),
        "Industry": p.get("industry"),
        "Country": p.get("country"),
        "Exchange": p.get("exchange"),
        "Price": price,
        "1D Chg": num(p.get("change")),
        "1D Chg %": num(p.get("changePercentage")),
        "52W Range": p.get("range"),
        "Market Cap": num(p.get("marketCap")),
        "Beta": num(p.get("beta")),
        "Volume": num(p.get("volume")),
        "Avg Volume": num(p.get("averageVolume")),
        "Forecast 12M (Target)": tcons,
        "Target High": num(t.get("targetHigh")),
        "Target Low": num(t.get("targetLow")),
        "Target Median": num(t.get("targetMedian")),
        "Upside to Target %": upside,
        "Analysts": int(total) if total else None,
        "Strong Buy": int(sb) if total else None,
        "Buy": int(b) if total else None,
        "Hold": int(h) if total else None,
        "Sell": int(s) if total else None,
        "Strong Sell": int(ss) if total else None,
        "% Buy": pct_buy,
        "% Strong Buy": pct_sbuy,
        "Consensus": g.get("consensus"),
        "P/E": num(r.get("priceToEarningsRatioTTM")),
        "PEG": num(r.get("priceToEarningsGrowthRatioTTM")),
        "P/B": num(r.get("priceToBookRatioTTM")),
        "P/S": num(r.get("priceToSalesRatioTTM")),
        "P/FCF": num(r.get("priceToFreeCashFlowRatioTTM")),
        "Dividend Yield %": pct(r.get("dividendYieldTTM")),
        "Div/Share": num(r.get("dividendPerShareTTM")) if r.get("dividendPerShareTTM") is not None else num(p.get("lastDividend")),
        "Payout %": pct(r.get("dividendPayoutRatioTTM")),
        "EPS (TTM)": num(r.get("netIncomePerShareTTM")),
        "Rev/Share": num(r.get("revenuePerShareTTM")),
        "Gross Margin %": pct(r.get("grossProfitMarginTTM")),
        "Oper Margin %": pct(r.get("operatingProfitMarginTTM")),
        "Net Margin %": pct(r.get("netProfitMarginTTM")),
        "Debt/Equity": num(r.get("debtToEquityRatioTTM")),
        "Current Ratio": num(r.get("currentRatioTTM")),
        "Employees": num(p.get("fullTimeEmployees")),
        "CEO": p.get("ceo"),
        "IPO Date": p.get("ipoDate"),
    }


COLUMNS = list(build_row("X", {}).keys())

# Number formats per column
FMT = {
    "Price": "#,##0.00", "1D Chg": "#,##0.00", "1D Chg %": "0.00",
    "Market Cap": "#,##0", "Beta": "0.00", "Volume": "#,##0", "Avg Volume": "#,##0",
    "Forecast 12M (Target)": "#,##0.00", "Target High": "#,##0.00", "Target Low": "#,##0.00",
    "Target Median": "#,##0.00", "Upside to Target %": "0.0", "% Buy": "0.0",
    "% Strong Buy": "0.0", "P/E": "0.00", "PEG": "0.00", "P/B": "0.00", "P/S": "0.00",
    "P/FCF": "0.00", "Dividend Yield %": "0.00", "Div/Share": "0.00", "Payout %": "0.0",
    "EPS (TTM)": "0.00", "Rev/Share": "0.00", "Gross Margin %": "0.0", "Oper Margin %": "0.0",
    "Net Margin %": "0.0", "Debt/Equity": "0.00", "Current Ratio": "0.00", "Employees": "#,##0",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SP_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_sheet(ws, rows, title_note=None):
    ws.append(COLUMNS)
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for row in rows:
        ws.append([row.get(col) for col in COLUMNS])

    # formats + light borders + highlight S&P rows
    sp_col = COLUMNS.index("In S&P 500") + 1
    for ridx in range(2, ws.max_row + 1):
        is_sp = ws.cell(row=ridx, column=sp_col).value == "Yes"
        for cidx, col in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=ridx, column=cidx)
            if col in FMT:
                cell.number_format = FMT[col]
            if is_sp:
                cell.fill = SP_FILL
    # column widths
    for cidx, col in enumerate(COLUMNS, start=1):
        width = max(10, min(34, len(col) + 2))
        if col in ("Company", "Industry", "CEO"):
            width = 30
        ws.column_dimensions[get_column_letter(cidx)].width = width
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    # conditional color scales on key columns
    for colname in ("Upside to Target %", "% Buy", "% Strong Buy"):
        ci = COLUMNS.index(colname) + 1
        L = get_column_letter(ci)
        rng = f"{L}2:{L}{ws.max_row}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="min", start_color="F8696B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="63BE7B"))


def summary_sheet(ws, rows):
    from collections import defaultdict
    agg = defaultdict(lambda: {"n": 0, "pe": [], "up": [], "buy": [], "mc": 0.0})
    for r in rows:
        sec = r.get("Sector") or "Unknown"
        a = agg[sec]
        a["n"] += 1
        if isinstance(r.get("P/E"), float) and 0 < r["P/E"] < 200:
            a["pe"].append(r["P/E"])
        if isinstance(r.get("Upside to Target %"), float):
            a["up"].append(r["Upside to Target %"])
        if isinstance(r.get("% Buy"), float):
            a["buy"].append(r["% Buy"])
        if isinstance(r.get("Market Cap"), float):
            a["mc"] += r["Market Cap"]

    def avg(x):
        return sum(x) / len(x) if x else None

    hdr = ["Sector", "# Stocks", "Total Market Cap", "Avg P/E", "Avg Upside %", "Avg % Buy"]
    ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for sec in sorted(agg, key=lambda s: -agg[s]["n"]):
        a = agg[sec]
        ws.append([sec, a["n"], a["mc"], avg(a["pe"]), avg(a["up"]), avg(a["buy"])])
    for ridx in range(2, ws.max_row + 1):
        ws.cell(row=ridx, column=3).number_format = "#,##0"
        for cc in (4, 5, 6):
            ws.cell(row=ridx, column=cc).number_format = "0.00"
    widths = [26, 10, 22, 12, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def about_sheet(ws, n_all, n_sp):
    lines = [
        ("Stock Analysis Workbook", ""),
        ("Generated", str(date.today())),
        ("Data source", "Financial Modeling Prep (real-time / TTM, via MCP)"),
        ("Total stocks", n_all),
        ("S&P 500 constituents included", n_sp),
        ("Parameters per stock", len(COLUMNS)),
        ("", ""),
        ("Notes", ""),
        ("S&P 500 rows are highlighted in pale yellow on every sheet.", ""),
        ("'Forecast 12M (Target)' is the consensus analyst 12-month price target.", ""),
        ("'Upside to Target %' = consensus target / price - 1.", ""),
        ("'% Buy' = (Strong Buy + Buy) / total analyst ratings.", ""),
        ("Ratios (P/E, margins, etc.) are trailing-twelve-month (TTM).", ""),
        ("Blank cells = value not provided by the data source for that ticker.", ""),
        ("", ""),
        ("Column definitions", ""),
    ]
    for k, v in lines:
        ws.append([k, v])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14, color="1F4E78")
    for r in range(2, 7):
        ws.cell(row=r, column=1).font = Font(bold=True)
    defs = [
        ("Beta", "Volatility vs market (1.0 = market)"),
        ("PEG", "P/E divided by earnings growth"),
        ("P/B / P/S / P/FCF", "Price to book / sales / free cash flow"),
        ("Debt/Equity", "Total debt relative to shareholder equity"),
        ("Current Ratio", "Current assets / current liabilities"),
        ("Consensus", "Overall analyst rating label"),
    ]
    for k, v in defs:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60


def main():
    recs = load_records()
    rows = [build_row(sym, obj) for sym, obj in recs.items()]
    # sort: S&P first, then by market cap desc
    rows.sort(key=lambda r: (0 if r["In S&P 500"] == "Yes" else 1,
                             -(r["Market Cap"] or 0)))
    n_sp = sum(1 for r in rows if r["In S&P 500"] == "Yes")

    wb = Workbook()
    about_sheet(wb.active, len(rows), n_sp)
    wb.active.title = "About"

    write_sheet(wb.create_sheet("All Stocks"), rows)
    write_sheet(wb.create_sheet("S&P 500"), [r for r in rows if r["In S&P 500"] == "Yes"])

    top = [r for r in rows if isinstance(r.get("% Buy"), float) and (r.get("Analysts") or 0) >= 5]
    top.sort(key=lambda r: (-(r["% Buy"]), -(r.get("Upside to Target %") or -999)))
    write_sheet(wb.create_sheet("Top Analyst Buys"), top[:200])

    summary_sheet(wb.create_sheet("Sector Summary"), rows)

    wb.save(XLSX)
    print(f"Wrote {XLSX}")
    print(f"  rows: {len(rows)} | S&P 500: {n_sp} | columns: {len(COLUMNS)}")
    filled = sum(1 for r in rows if isinstance(r.get("Price"), float))
    pe = sum(1 for r in rows if isinstance(r.get("P/E"), float))
    fc = sum(1 for r in rows if isinstance(r.get("Forecast 12M (Target)"), float))
    bu = sum(1 for r in rows if isinstance(r.get("% Buy"), float))
    print(f"  with price: {filled} | with P/E: {pe} | with forecast: {fc} | with buy%: {bu}")


if __name__ == "__main__":
    main()
